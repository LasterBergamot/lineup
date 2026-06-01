import io

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

# source: https://python-docx.readthedocs.io/en/latest/dev/analysis/features/text/tab-stops.html
from docx.enum.text import WD_TAB_LEADER, WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
from docx.table import _Cell


def should_keep_tab_stop(tab_stop):
    return tab_stop.leader == WD_TAB_LEADER.SPACES


def update_cell_text(cell, text):
    cell.text = text


def align_cell_vertically(cell, alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER):
    cell.vertical_alignment = alignment


def align_cell_horizontally(cell, alignment=WD_PARAGRAPH_ALIGNMENT.CENTER):
    cell.paragraphs[0].alignment = alignment


class DocumentManager:
    def __init__(self, filename):
        self.doc = Document(filename)

    def setup_style(self, style_name="Normal", font_name="Calibri", font_size=Pt(10)):
        style = self.doc.styles[style_name]
        font = style.font
        font.name = font_name
        font.size = font_size

    def remove_non_space_tab_stops(self):
        for paragraph in self.doc.paragraphs:
            paragraph_format = paragraph.paragraph_format
            tabs_to_keep = [
                (ts.position, ts.leader)
                for ts in paragraph_format.tab_stops
                if should_keep_tab_stop(ts)
            ]
            paragraph_format.tab_stops.clear_all()
            for position, leader in tabs_to_keep:
                paragraph_format.tab_stops.add_tab_stop(position, leader)

    def update_paragraph_text(self, index, text):
        if index < len(self.doc.paragraphs):
            paragraph = self.doc.paragraphs[index]
            paragraph.text = text
        else:
            print(f"Paragraph index {index} out of range.")

    def add_run(self, index, text):
        if index < len(self.doc.paragraphs):
            paragraph = self.doc.paragraphs[index]
            paragraph.add_run(text)
        else:
            print(f"Paragraph index {index} out of range.")

    def underline_run(self, index, text):
        if index < len(self.doc.paragraphs):
            paragraph = self.doc.paragraphs[index]
            for run in paragraph.runs:
                if run.text == text:
                    run.underline = True
        else:
            print(f"Paragraph index {index} out of range.")

    def make_run_bold(self, index, text):
        if index < len(self.doc.paragraphs):
            paragraph = self.doc.paragraphs[index]
            for run in paragraph.runs:
                if run.text == text:
                    run.bold = True
        else:
            print(f"Paragraph index {index} out of range.")

    def get_table_cell(self, table_index=0, row=1, column=1) -> _Cell:
        if table_index < len(self.doc.tables):
            table = self.doc.tables[table_index]
            if row < len(table.rows) and column < len(table.columns):
                return table.cell(row, column)
            else:
                print(f"Table cell ({row}, {column}) out of range.")
        else:
            print(f"Table index {table_index} out of range.")

    def save(self, filename):
        self.doc.save(filename)

    def to_bytes(self) -> bytes:
        buffer = io.BytesIO()
        self.doc.save(buffer)
        buffer.seek(0)
        return buffer.read()
