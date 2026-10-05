"""Low-level edits on the lineup `.docx` template with python-docx."""

import io

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

# source: https://python-docx.readthedocs.io/en/latest/dev/analysis/features/text/tab-stops.html
from docx.enum.text import WD_TAB_LEADER, WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
from docx.table import _Cell


def should_keep_tab_stop(tab_stop):
    """True for a tab stop whose leader is plain spaces (see `remove_non_space_tab_stops`)."""
    return tab_stop.leader == WD_TAB_LEADER.SPACES


def update_cell_text(cell, text):
    """Replace the text of a table cell."""
    cell.text = text


def align_cell_vertically(cell, alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER):
    """Set a cell's vertical alignment (centred by default)."""
    cell.vertical_alignment = alignment


def align_cell_horizontally(cell, alignment=WD_PARAGRAPH_ALIGNMENT.CENTER):
    """Set the alignment of a cell's first paragraph (centred by default)."""
    cell.paragraphs[0].alignment = alignment


class DocumentManager:
    """Thin wrapper over python-docx for editing the lineup template.

    Paragraphs and table cells are addressed by position, so it depends on the layout of
    `resources/rajtlista.docx`.
    """

    def __init__(self, filename):
        """Open the `.docx` at `filename` (a path or file-like object)."""
        self.doc = Document(filename)

    def setup_style(self, style_name="Normal", font_name="Calibri", font_size=Pt(10)):
        """Set the font name and size of a paragraph style (the template's `Normal` by default)."""
        style = self.doc.styles[style_name]
        font = style.font
        font.name = font_name
        font.size = font_size

    def remove_non_space_tab_stops(self):
        """Drop every tab stop whose leader is not plain spaces, in every paragraph.

        Leaders such as dots or lines are filler characters drawn up to the tab stop, which
        would show between the label and the value we insert; space-leader stops are kept.
        """
        for paragraph in self.doc.paragraphs:
            paragraph_format = paragraph.paragraph_format
            tabs_to_keep = [
                (ts.position, ts.leader)
                for ts in paragraph_format.tab_stops
                if should_keep_tab_stop(ts)
            ]
            paragraph_format.tab_stops.clear_all()
            for position, leader in tabs_to_keep:
                paragraph_format.tab_stops.add_tab_stop(position, leader)

    def update_paragraph_text(self, index, text):
        """Replace a paragraph's text by index; an out-of-range index is reported with `print` and ignored."""
        if index < len(self.doc.paragraphs):
            paragraph = self.doc.paragraphs[index]
            paragraph.text = text
        else:
            print(f"Paragraph index {index} out of range.")

    def add_run(self, index, text):
        """Append a text run to a paragraph, so it can be styled separately. Out-of-range indexes are ignored."""
        if index < len(self.doc.paragraphs):
            paragraph = self.doc.paragraphs[index]
            paragraph.add_run(text)
        else:
            print(f"Paragraph index {index} out of range.")

    def underline_run(self, index, text):
        """Underline the runs of a paragraph whose text equals `text`."""
        if index < len(self.doc.paragraphs):
            paragraph = self.doc.paragraphs[index]
            for run in paragraph.runs:
                if run.text == text:
                    run.underline = True
        else:
            print(f"Paragraph index {index} out of range.")

    def make_run_bold(self, index, text):
        """Make the runs of a paragraph whose text equals `text` bold."""
        if index < len(self.doc.paragraphs):
            paragraph = self.doc.paragraphs[index]
            for run in paragraph.runs:
                if run.text == text:
                    run.bold = True
        else:
            print(f"Paragraph index {index} out of range.")

    def get_table_cell(self, table_index=0, row=1, column=1) -> _Cell:
        """Return a table cell, or `None` (after printing a message) if the table, row or column is out of range."""
        if table_index < len(self.doc.tables):
            table = self.doc.tables[table_index]
            if row < len(table.rows) and column < len(table.columns):
                return table.cell(row, column)
            else:
                print(f"Table cell ({row}, {column}) out of range.")
        else:
            print(f"Table index {table_index} out of range.")

    def save(self, filename):
        """Write the document to a file path."""
        self.doc.save(filename)

    def to_bytes(self) -> bytes:
        """Serialise the document to `.docx` bytes in memory (nothing touches the disk)."""
        buffer = io.BytesIO()
        self.doc.save(buffer)
        buffer.seek(0)
        return buffer.read()
